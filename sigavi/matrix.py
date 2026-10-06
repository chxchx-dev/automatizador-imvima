from __future__ import annotations

import posixpath
import re
import shutil
import warnings
from copy import copy
from datetime import date, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .constants import MATRIX_HEADERS, SHEET_NAME
from .duplicates import normalize_text, normalize_url
from .timeutil import today_colombia


class MatrixError(RuntimeError):
    pass


def _worksheet_part(path: Path, sheet_title: str) -> str | None:
    workbook_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    try:
        with ZipFile(path, "r") as archive:
            workbook_xml = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    except (KeyError, OSError, ValueError):
        return None
    targets = {item.attrib.get("Id"): item.attrib.get("Target") for item in relationships.findall(f"{{{package_rel_ns}}}Relationship")}
    sheets = workbook_xml.find(f"{{{workbook_ns}}}sheets")
    if sheets is None:
        return None
    for item in sheets.findall(f"{{{workbook_ns}}}sheet"):
        if item.attrib.get("name") != sheet_title:
            continue
        target = targets.get(item.attrib.get(f"{{{rel_ns}}}id"))
        if not target:
            return None
        return posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target))
    return None


def _preserve_template_extensions(
    template_path: Path,
    output_path: Path,
    template_worksheet_part: str,
    output_worksheet_part: str,
    last_data_row: int,
) -> None:
    """Restore worksheet OOXML extensions that openpyxl does not understand."""
    try:
        with ZipFile(template_path, "r") as source_zip:
            source_xml = source_zip.read(template_worksheet_part).decode("utf-8")
    except (KeyError, OSError, UnicodeDecodeError, ValueError):
        return

    extension_match = re.search(r"<extLst\b.*?</extLst>", source_xml, flags=re.DOTALL)
    if not extension_match:
        return
    extension_xml = extension_match.group(0)
    extension_xml = re.sub(
        r"(<xm:sqref>G\d+:G)\d+(</xm:sqref>)",
        rf"\g<1>{last_data_row}\g<2>",
        extension_xml,
    )

    temporary_path: Path | None = None
    with NamedTemporaryFile(prefix="sigavi-ooxml-", suffix=".xlsx", dir=output_path.parent, delete=False) as handle:
        temporary_path = Path(handle.name)
    try:
        with ZipFile(output_path, "r") as current_zip, ZipFile(temporary_path, "w", ZIP_DEFLATED) as patched_zip:
            for member in current_zip.infolist():
                payload = current_zip.read(member.filename)
                if member.filename == output_worksheet_part:
                    worksheet_xml = payload.decode("utf-8")
                    if "<extLst" not in worksheet_xml:
                        if "xr:uid" in extension_xml and "xmlns:xr=" not in worksheet_xml:
                            worksheet_xml = re.sub(
                                r"<worksheet\b",
                                '<worksheet xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision"',
                                worksheet_xml,
                                count=1,
                            )
                        worksheet_xml = worksheet_xml.replace("</worksheet>", extension_xml + "</worksheet>")
                    payload = worksheet_xml.encode("utf-8")
                patched_zip.writestr(member, payload)
        temporary_path.replace(output_path)
    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise


def _date_value(value: Any, epoch) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (float, int)):
        try:
            converted = from_excel(value, epoch)
            return converted.date() if isinstance(converted, datetime) else converted
        except (ValueError, OverflowError, TypeError):
            return None
    if isinstance(value, str):
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return None
    return None


def _row_is_alert(row: tuple[Any, ...]) -> bool:
    return any(value not in (None, "") for value in row[:17])


def _alert_key(code: Any, alert_date: Any, product: Any, link: Any, epoch) -> tuple[str, str, str, str]:
    normalized_date = _date_value(alert_date, epoch)
    return (
        normalize_text(str(code or "")),
        normalized_date.isoformat() if normalized_date else "",
        normalize_text(str(product or "")),
        normalize_url(str(link or "")),
    )


def _record_values(record: dict[str, Any]) -> list[Any]:
    alert_date = date.fromisoformat(record["alert_date"]) if record.get("alert_date") else None
    review_date = today_colombia()
    return [
        review_date,
        alert_date,
        record.get("code") or None,
        "INVIMA",
        record.get("pdf_url") or record.get("alert_url") or None,
        record.get("document_type") or None,
        record.get("classification") or None,
        record.get("product_name") or None,
        record.get("registry") or None,
        record.get("holder") or None,
        record.get("manufacturer_importer") or None,
        record.get("reference_code") or None,
        record.get("lot_serial") or None,
        record.get("description") or None,
        record.get("indication_use") or None,
        None,
        None,
        None,
    ]


class MatrixWriter:
    def __init__(self, template_path: str | Path, output_path: str | Path) -> None:
        self.template_path = Path(template_path)
        self.output_path = Path(output_path)

    def sync(self, records: list[dict[str, Any]]) -> tuple[dict[int, int], int, int]:
        if not self.template_path.is_file():
            raise MatrixError(f"No se encontró la plantilla institucional: {self.template_path}")
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        if self.template_path.resolve() == self.output_path.resolve():
            raise MatrixError("La matriz de salida no puede sobrescribir la plantilla original.")
        if not self.output_path.exists():
            shutil.copy2(self.template_path, self.output_path)

        try:
            with warnings.catch_warnings():
                warnings.filterwarnings(
                    "ignore",
                    message="Data Validation extension is not supported and will be removed",
                    category=UserWarning,
                )
                workbook = load_workbook(self.output_path, keep_links=True)
        except Exception as exc:
            raise MatrixError(f"No se pudo abrir el libro Excel: {exc}") from exc
        sheet = next((ws for ws in workbook.worksheets if normalize_text(ws.title) == normalize_text(SHEET_NAME)), None)
        if sheet is None:
            workbook.close()
            raise MatrixError(f"La plantilla no contiene la hoja «{SHEET_NAME}».")
        template_worksheet_part = _worksheet_part(self.template_path, SHEET_NAME)
        if not template_worksheet_part:
            workbook.close()
            raise MatrixError("No se pudo identificar la hoja objetivo dentro del archivo XLSX.")

        headers = tuple(str(sheet.cell(5, col).value or "").strip() for col in range(1, 19))
        if headers != MATRIX_HEADERS:
            workbook.close()
            raise MatrixError("Los encabezados de la fila 5 no coinciden exactamente con las 18 columnas institucionales.")

        filter_end = 164
        if sheet.auto_filter.ref and ":" in sheet.auto_filter.ref:
            try:
                filter_end = int(sheet.auto_filter.ref.split(":")[-1].lstrip("$ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
            except ValueError:
                pass
        body_start = 6
        body_end = max(filter_end, body_start)

        existing_rows: dict[tuple[str, str, str, str], int] = {}
        existing_composite_rows: dict[tuple[str, str, str], list[tuple[str, int]]] = {}
        for row_number in range(body_start, sheet.max_row + 1):
            row_values = tuple(sheet.cell(row_number, col).value for col in range(1, 19))
            if not _row_is_alert(row_values):
                continue
            key = _alert_key(row_values[2], row_values[1], row_values[7], row_values[4], workbook.epoch)
            if any(key):
                existing_rows.setdefault(key, row_number)
                if key[0] and key[1] and key[3]:
                    existing_composite_rows.setdefault((key[0], key[1], key[3]), []).append((key[2], row_number))

        to_add: list[dict[str, Any]] = []
        ids_to_rows: dict[int, int] = {}
        links_repaired = False
        queued_keys: set[tuple[str, str, str, str]] = set()
        queued_aliases: dict[tuple[str, str, str, str], list[int]] = {}
        skipped = 0
        for record in records:
            key = _alert_key(record.get("code"), record.get("alert_date"), record.get("product_name"), record.get("pdf_url") or record.get("alert_url"), workbook.epoch)
            composite = (key[0], key[1], key[3])
            existing_row = existing_rows.get(key)
            if not existing_row:
                for existing_product, row_number in existing_composite_rows.get(composite, []):
                    if not existing_product or not key[2] or existing_product == key[2]:
                        existing_row = row_number
                        break
            if existing_row:
                skipped += 1
                # Older matrix copies may have the URL text but lack an actual
                # Excel hyperlink. Repair both the displayed value and target.
                pdf_url = record.get("pdf_url") or record.get("alert_url")
                if pdf_url:
                    link_cell = sheet.cell(existing_row, 5)
                    if link_cell.value in (None, ""):
                        link_cell.value = pdf_url
                        links_repaired = True
                    if not link_cell.hyperlink or link_cell.hyperlink.target != pdf_url:
                        link_cell.hyperlink = pdf_url
                        links_repaired = True
                if record.get("id") is not None:
                    ids_to_rows[int(record["id"])] = existing_row
                continue
            if key in queued_keys:
                skipped += 1
                if record.get("id") is not None:
                    queued_aliases.setdefault(key, []).append(int(record["id"]))
                continue
            queued_keys.add(key)
            to_add.append(record)

        if not to_add and not links_repaired:
            workbook.close()
            output_worksheet_part = _worksheet_part(self.output_path, SHEET_NAME)
            if output_worksheet_part:
                try:
                    last_data_row = max([body_end, *ids_to_rows.values()])
                    _preserve_template_extensions(
                        self.template_path,
                        self.output_path,
                        template_worksheet_part,
                        output_worksheet_part,
                        last_data_row,
                    )
                except Exception as exc:
                    raise MatrixError(f"No se pudieron conservar las validaciones especiales de Excel: {exc}") from exc
            return ids_to_rows, 0, skipped

        if not to_add:
            # Save repairs for rows already present in the destination workbook.
            with NamedTemporaryFile(prefix="sigavi-matrix-", suffix=".xlsx", dir=self.output_path.parent, delete=False) as handle:
                temp_path = Path(handle.name)
            try:
                workbook.save(temp_path)
                workbook.close()
                output_worksheet_part = _worksheet_part(temp_path, SHEET_NAME)
                if not output_worksheet_part:
                    raise MatrixError("No se pudo identificar la hoja objetivo en la copia XLSX guardada.")
                _preserve_template_extensions(
                    self.template_path,
                    temp_path,
                    template_worksheet_part,
                    output_worksheet_part,
                    max([body_end, *ids_to_rows.values()]),
                )
                temp_path.replace(self.output_path)
            except Exception:
                workbook.close()
                temp_path.unlink(missing_ok=True)
                raise
            return ids_to_rows, 0, skipped

        # Reuse the workbook's preformatted blank rows first; rows with only the
        # institution's prefilled responsible-person value remain available.
        blank_rows = []
        for row_number in range(body_start, body_end + 1):
            values = tuple(sheet.cell(row_number, col).value for col in range(1, 18))
            if not any(value not in (None, "") for value in values):
                blank_rows.append(row_number)

        if len(blank_rows) < len(to_add):
            overflow = len(to_add) - len(blank_rows)
            insert_at = body_end + 1
            sheet.insert_rows(insert_at, overflow)
            style_row = min(max(body_start + 1, 7), sheet.max_row - overflow)
            for offset in range(overflow):
                target_row = insert_at + offset
                for col in range(1, 19):
                    source = sheet.cell(style_row, col)
                    target = sheet.cell(target_row, col)
                    if source.has_style:
                        target._style = copy(source._style)
                    if source.number_format:
                        target.number_format = source.number_format
                    if source.alignment:
                        target.alignment = copy(source.alignment)
                    if col == 18 and source.value not in (None, ""):
                        target.value = source.value
                if sheet.row_dimensions[style_row].height:
                    sheet.row_dimensions[target_row].height = sheet.row_dimensions[style_row].height
            blank_rows.extend(range(insert_at, insert_at + overflow))

        for record, row_number in zip(to_add, blank_rows):
            values = _record_values(record)
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row_number, col)
                if col == 18:
                    # Keep the template's institutional default, if it has one.
                    continue
                cell.value = value
            link_cell = sheet.cell(row_number, 5)
            pdf_url = record.get("pdf_url") or record.get("alert_url")
            if pdf_url:
                link_cell.hyperlink = pdf_url
            record_id = int(record["id"])
            ids_to_rows[record_id] = row_number
            key = _alert_key(record.get("code"), record.get("alert_date"), record.get("product_name"), record.get("pdf_url") or record.get("alert_url"), workbook.epoch)
            for alias_id in queued_aliases.get(key, []):
                ids_to_rows[alias_id] = row_number

        new_end = max(body_end, max(ids_to_rows.values()))
        sheet.auto_filter.ref = f"A5:R{new_end}"
        for validation in sheet.data_validations.dataValidation:
            if validation.type == "list" and validation.sqref:
                ranges = []
                for item in str(validation.sqref).split():
                    if item.startswith("F") and ":F" in item:
                        start = item.split(":", 1)[0]
                        ranges.append(f"{start}:F{new_end}")
                    else:
                        ranges.append(item)
                validation.sqref = " ".join(ranges)

        if sheet.print_area:
            try:
                area = str(sheet.print_area)
                if ":" in area:
                    first, last = area.split(":", 1)
                    last_col = "".join(ch for ch in last if ch.isalpha()) or "R"
                    last_row = int("".join(ch for ch in last if ch.isdigit()) or "0")
                    if new_end > last_row:
                        sheet.print_area = f"{first}:{last_col}{new_end + max(0, last_row - body_end)}"
            except (TypeError, ValueError):
                pass

        with NamedTemporaryFile(prefix="sigavi-matrix-", suffix=".xlsx", dir=self.output_path.parent, delete=False) as handle:
            temp_path = Path(handle.name)
        try:
            workbook.save(temp_path)
            workbook.close()
            output_worksheet_part = _worksheet_part(temp_path, SHEET_NAME)
            if not output_worksheet_part:
                raise MatrixError("No se pudo identificar la hoja objetivo en la copia XLSX guardada.")
            _preserve_template_extensions(
                self.template_path,
                temp_path,
                template_worksheet_part,
                output_worksheet_part,
                new_end,
            )
            temp_path.replace(self.output_path)
        except Exception:
            workbook.close()
            temp_path.unlink(missing_ok=True)
            raise
        return ids_to_rows, len(to_add), skipped
